"""Upload validation, text extraction (.md .txt .pdf .docx) and the two chunking strategies.

OCR is NOT included: scanned/image-only PDFs are rejected with a clear error.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass

ALLOWED = {".md": "md", ".txt": "txt", ".pdf": "pdf", ".docx": "docx"}
TOKEN_RE = re.compile(r"\w+|[^\w\s]", re.UNICODE)  # "word tokens": words + punctuation marks
SENTENCE_RE = re.compile(r"(?<=[.!?])\s+")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


class IngestError(ValueError):
    pass


@dataclass
class Segment:
    """A contiguous piece of extracted text with its location in the original."""
    text: str
    page: int | None = None
    section: str | None = None


@dataclass
class Chunk:
    text: str
    page: int | None
    section: str | None
    start_char: int
    end_char: int
    n_chars: int
    n_tokens: int


def count_tokens(text: str) -> int:
    return len(TOKEN_RE.findall(text))


# ---------------- validation ----------------

def safe_filename(name: str) -> str:
    """Display-only name: strips any directory parts and unsafe characters (files are stored under random ids)."""
    name = re.split(r"[\\/]", name or "")[-1]
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")
    return name[:120] or "upload"


def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    """Returns the file type or raises IngestError. Checks extension AND content signature."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED:
        raise IngestError(f"Unsupported file type '{ext or 'none'}'. Allowed: {', '.join(ALLOWED)}")
    if not data:
        raise IngestError("File is empty")
    if len(data) > max_bytes:
        raise IngestError(f"File exceeds {max_bytes // (1024 * 1024)} MB limit")
    ftype = ALLOWED[ext]
    if ftype == "pdf" and not data.startswith(b"%PDF-"):
        raise IngestError("File content is not a PDF")
    if ftype == "docx":
        if not data.startswith(b"PK"):
            raise IngestError("File content is not a DOCX")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
                if "word/document.xml" not in names:
                    raise IngestError("DOCX is missing word/document.xml")
                if sum(i.file_size for i in z.infolist()) > max_bytes * 20:
                    raise IngestError("DOCX expands too much (possible zip bomb)")
                if any(n.lower().startswith("word/vbaproject") for n in names):
                    raise IngestError("Macro-enabled documents are not accepted")
        except zipfile.BadZipFile as exc:
            raise IngestError("Corrupt DOCX") from exc
    if ftype in ("md", "txt"):
        if b"\x00" in data[:4096]:
            raise IngestError("Binary content in a text file")
        try:
            data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise IngestError("Text files must be UTF-8") from exc
    return ftype


# ---------------- extraction ----------------

def extract(ftype: str, data: bytes) -> list[Segment]:
    if ftype in ("md", "txt"):
        return _extract_text(data.decode("utf-8"), markdown=(ftype == "md"))
    if ftype == "pdf":
        return _extract_pdf(data)
    if ftype == "docx":
        return _extract_docx(data)
    raise IngestError("Unsupported type")


def _extract_text(text: str, markdown: bool) -> list[Segment]:
    text = text.replace("\r\n", "\n")
    if not markdown:
        return [Segment(text)] if text.strip() else []
    segs: list[Segment] = []
    section, buf = None, []
    for line in text.split("\n"):
        m = HEADING_RE.match(line)
        if m:
            if "".join(buf).strip():
                segs.append(Segment("\n".join(buf).strip(), section=section))
            section, buf = m.group(2).strip(), [line]
        else:
            buf.append(line)
    if "".join(buf).strip():
        segs.append(Segment("\n".join(buf).strip(), section=section))
    return segs


def _extract_pdf(data: bytes) -> list[Segment]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise IngestError("Encrypted PDFs are not supported")
        segs = [Segment(t.strip(), page=i + 1) for i, p in enumerate(reader.pages) if (t := p.extract_text() or "").strip()]
    except IngestError:
        raise
    except Exception as exc:
        raise IngestError(f"Could not read PDF: {exc}") from exc
    if not segs:
        raise IngestError("PDF has no extractable text (scanned PDF? OCR is not included)")
    return segs


def _extract_docx(data: bytes) -> list[Segment]:
    import docx

    d = docx.Document(io.BytesIO(data))
    segs: list[Segment] = []
    section, buf = None, []
    for para in d.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            if buf:
                segs.append(Segment("\n\n".join(buf), section=section))
            section, buf = para.text.strip(), [f"# {para.text.strip()}"]
        elif para.text.strip():
            buf.append(para.text.strip())
    for table in d.tables:  # tables appended as pipe-separated rows
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
        buf.append("\n".join(rows))
    if buf:
        segs.append(Segment("\n\n".join(buf), section=section))
    if not any(s.text.strip() for s in segs):
        raise IngestError("DOCX contains no text")
    return segs


# ---------------- chunking ----------------

def _units(text: str, unit: str) -> list[tuple[int, int]]:
    """Spans of the measuring unit: each char, or each token (with trailing whitespace)."""
    if unit == "chars":
        return [(i, i + 1) for i in range(len(text))]
    spans = [m.span() for m in TOKEN_RE.finditer(text)]
    return [(s, spans[i + 1][0] if i + 1 < len(spans) else len(text)) for i, (s, _) in enumerate(spans)]


def _measure(text: str, unit: str) -> int:
    return len(text) if unit == "chars" else count_tokens(text)


def chunk_fixed(text: str, size: int, overlap: int, unit: str) -> list[tuple[int, int]]:
    """Sliding window of `size` units, stepping by size - overlap. Returns char spans."""
    if unit == "chars":
        step = size - overlap
        return [(i, min(i + size, len(text))) for i in range(0, max(len(text) - overlap, 1), step)]
    units = _units(text, unit)
    out, step, i = [], size - overlap, 0
    while i < len(units):
        j = min(i + size, len(units))
        out.append((units[i][0], units[j - 1][1]))
        if j == len(units):
            break
        i += step
    return out


def _split_boundaries(text: str, boundaries: list[str]) -> list[tuple[int, int]]:
    """Splits text into the smallest boundary pieces requested (heading > paragraph > sentence)."""
    pieces = [(0, len(text))]
    if "heading" in boundaries:
        starts = [0] + [m.start() for m in re.finditer(r"(?m)^#{1,6}\s", text) if m.start() > 0]
        pieces = [(s, e) for s, e in zip(starts, starts[1:] + [len(text)])]
    if "paragraph" in boundaries:
        pieces = [p for s, e in pieces for p in _regex_split(text, s, e, r"\n\s*\n")]
    if "sentence" in boundaries:
        pieces = [p for s, e in pieces for p in _regex_split(text, s, e, SENTENCE_RE.pattern)]
    return [(s, e) for s, e in pieces if text[s:e].strip()]


def _regex_split(text: str, start: int, end: int, pattern: str) -> list[tuple[int, int]]:
    out, pos = [], start
    for m in re.finditer(pattern, text[start:end]):
        out.append((pos, start + m.end()))
        pos = start + m.end()
    out.append((pos, end))
    return out


def chunk_recursive(text: str, size: int, overlap: int, unit: str, boundaries: list[str]) -> list[tuple[int, int]]:
    """Packs boundary pieces into chunks up to `size`; oversize pieces fall back to fixed windows.
    Overlap is carried as whole trailing pieces whose total size <= overlap."""
    pieces: list[tuple[int, int]] = []
    for s, e in _split_boundaries(text, boundaries or ["paragraph"]):
        if _measure(text[s:e], unit) > size:
            pieces += [(s + a, s + b) for a, b in chunk_fixed(text[s:e], size, min(overlap, size // 2), unit)]
        else:
            pieces.append((s, e))
    chunks, cur = [], []
    for piece in pieces:
        if cur and _measure(text[cur[0][0]:piece[1]], unit) > size:
            chunks.append((cur[0][0], cur[-1][1]))
            carry: list[tuple[int, int]] = []
            for p in reversed(cur):
                if overlap and _measure(text[p[0]:cur[-1][1]], unit) <= overlap:
                    carry.insert(0, p)
                else:
                    break
            cur = carry if carry and _measure(text[carry[0][0]:piece[1]], unit) <= size else []
        cur.append(piece)
    if cur:
        chunks.append((cur[0][0], cur[-1][1]))
    return chunks


def chunk_segments(segments: list[Segment], cfg: dict) -> list[Chunk]:
    out: list[Chunk] = []
    offset = 0
    for seg in segments:
        if cfg["strategy"] == "fixed":
            spans = chunk_fixed(seg.text, cfg["size"], cfg["overlap"], cfg["unit"])
        else:
            spans = chunk_recursive(seg.text, cfg["size"], cfg["overlap"], cfg["unit"], cfg.get("boundaries", []))
        for s, e in spans:
            raw = seg.text[s:e]
            body = raw.strip()
            if not body:
                continue
            lead = len(raw) - len(raw.lstrip())
            out.append(Chunk(body, seg.page, seg.section, offset + s + lead, offset + s + lead + len(body),
                             len(body), count_tokens(body)))
        offset += len(seg.text) + 2
    return out
